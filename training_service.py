"""Training & onboarding content: PowerPoint decks + orientation quiz.

Decks built here with python-pptx:
- ``worker`` deck: how to use the phone app and website, expected behavior in hosts'
  homes, ethics, sexual harassment policy, and other new-hire orientation items.
- ``host`` deck: how Broom Service works, services & pricing, the host portal, and
  funding/capital partners for leads.
- ``construction`` deck: Buildstack Construction crew orientation — OSHA-10 baseline
  (fatal four), PPE, ladder/scaffold, hazard communication, lockout/tagout, silica/dust/
  asbestos awareness, pay app & direct deposit, harassment policy, and a trades knowledge
  test.

Each deck is returned as in-memory ``.pptx`` bytes for storage in the document
library. The Broom worker quiz (10 questions) and the construction OSHA-10 / trades
quiz (``OSHA_QUIZ``) are defined here too.
"""

QUIZ = [
    {
        "q": "How do you get the worker app on your phone?",
        "options": [
            "Download it from the App Store",
            "Open the site on your phone and save it to your home screen",
            "It installs itself",
        ],
        "answer": 1,
    },
    {
        "q": "Why do you clock in and clock out on every job?",
        "options": [
            "So the office knows where you are",
            "Because that's how you get paid",
            "Just for fun",
        ],
        "answer": 1,
    },
    {
        "q": "When you arrive at a property, what should you do first?",
        "options": [
            "Clock in when you're inside the work area and then start cleaning",
            "Text your friends",
            "Wait outside until someone tells you to start",
        ],
        "answer": 0,
    },
    {
        "q": "Where do you find directions to your job?",
        "options": [
            "Call the office every time",
            "The map in the worker app",
            "Guess the address",
        ],
        "answer": 1,
    },
    {
        "q": "Can you eat food, use the bed, or take supplies from the host's home?",
        "options": [
            "Yes, if you're careful",
            "Only snacks",
            "No — it's the host's home; you treat everything with respect",
        ],
        "answer": 2,
    },
    {
        "q": "What should you do if you see damage or something broken at a property?",
        "options": [
            "Hide it",
            "Report it right away with a photo",
            "Fix it with tape",
        ],
        "answer": 1,
    },
    {
        "q": "How should you treat every person you meet at a property or office?",
        "options": [
            "With respect, always",
            "Only if they talk to you first",
            "Whatever, it doesn't matter",
        ],
        "answer": 0,
    },
    {
        "q": "Is unwelcome touching or sexual comments ever okay at work?",
        "options": [
            "No, never — it's harassment and won't be tolerated",
            "Yes if someone jokes about it",
            "Sometimes",
        ],
        "answer": 0,
    },
    {
        "q": "Where can you report a problem or concern safely?",
        "options": [
            "Nowhere",
            "Tell the owner or use the reporting channel — it's confidential",
            "Post it on social media",
        ],
        "answer": 1,
    },
    {
        "q": "Where do you find your paycheck and pay stubs?",
        "options": [
            "Under 'Pay' in the worker app or portal",
            "The host sends it by mail",
            "You have to ask in person",
        ],
        "answer": 0,
    },
]


OSHA_QUIZ = [
    {
        "q": "The 'Fatal Four' on construction sites are falls, struck-by, caught-in/between, and ______.",
        "options": ["Electrical", "Sun exposure", "Fatigue", "Loud noise"],
        "answer": 0,
        "topic": "OSHA-10 fatal four",
    },
    {
        "q": "At what height does fall protection (guardrails, harness + anchor) become required on most residential/commercial construction?",
        "options": [
            "3 feet",
            "6 feet",
            "15 feet",
            "Only on roofs",
        ],
        "answer": 1,
        "topic": "Fall protection",
    },
    {
        "q": "When working near or over moving equipment (forklift, excavator), what's the key struck-by protection?",
        "options": [
            "Stay in the equipment operator's blind spot",
            "Stay out of the swing/reach path and wear hi-vis so the operator sees you",
            "Wave your arms so they know to stop",
            "There's no way to prevent struck-by injuries",
        ],
        "answer": 1,
        "topic": "Struck-by",
    },
    {
        "q": "Which PPE is required on every active job site?",
        "options": [
            "Hard hat + eye protection",
            "Long sleeves only",
            "Steel-toe boots, nothing else",
            "PPE is optional for experienced crew",
        ],
        "answer": 0,
        "topic": "PPE",
    },
    {
        "q": "On a ladder, what is the correct body position?",
        "options": [
            "Lean backward while reaching",
            "Keep 3 points of contact and stay below the top rung",
            "Stand on the very top to reach high",
            "Face away from the ladder",
        ],
        "answer": 1,
        "topic": "Ladder safety",
    },
    {
        "q": "What does Lockout/Tagout (LOTO) prevent?",
        "options": [
            "Theft of tools",
            "Accidental energy release when working on electrical, plumbing, or equipment",
            "Getting paint on uniforms",
            "Nothing — it's just paperwork",
        ],
        "answer": 1,
        "topic": "Lockout/tagout",
    },
    {
        "q": "On a tape measure, the smallest marked increment is usually:",
        "options": ["1 inch", "1/2 inch", "1/16 inch", "2 inches"],
        "answer": 2,
        "topic": "Reading a tape measure",
    },
    {
        "q": "A 12-ft wall at 8 ft tall needs how many square feet of drywall?",
        "options": ["64 sq ft", "96 sq ft", "80 sq ft", "120 sq ft"],
        "answer": 1,
        "topic": "Simple trade math",
    },
    {
        "q": "If you work 46 hours in one week at a regular hourly rate, which hours are overtime?",
        "options": [
            "None",
            "Hours 41-46 (6 hours at 1.5x)",
            "The first 40, then rebuild the schedule",
            "Only Saturday hours",
        ],
        "answer": 1,
        "topic": "Payroll & Overtime",
    },
    {
        "q": "To control silica dust from cutting concrete or masonry, the best practice is:",
        "options": [
            "Wet cutting + dust collection + a respirator when required",
            "Just hold your breath",
            "Wear a bandana",
            "Cut faster so dust spreads less",
        ],
        "answer": 0,
        "topic": "Silica / dust awareness",
    },
    {
        "q": "Unwelcome touching, comments, or sexual jokes on the job site are:",
        "options": [
            "Fine if everyone laughs",
            "Harassment — zero tolerance, report confidentially, no retaliation",
            "Only a problem in the office",
            "Allowed between friends",
        ],
        "answer": 1,
        "topic": "Workplace harassment policy",
    },
    {
        "q": "Before working near electrical panels or exposed wires you should:",
        "options": [
            "Assume it's dead and start cutting",
            "LOTO it, verify with a tester, and use GFCI where needed",
            "Just avoid touching metal",
            "Stand on a metal ladder",
        ],
        "answer": 1,
        "topic": "Electrical safety",
    },
]


BRAND = {
    "title_color": (49, 46, 129),  # indigo-900
    "accent_color": (99, 102, 241),  # indigo-500
    "dark": (15, 23, 42),  # slate-900
    "light": (248, 250, 252),  # slate-50
}


def _add_title_slide(prs, doc_title, subtitle):
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    title = slide.shapes.title
    title.text = doc_title
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
    title.text_frame.paragraphs[0].runs[0].font.bold = True
    title.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(*BRAND["title_color"])
    if slide.shapes.placeholders and len(slide.shapes.placeholders) > 1:
        sub = slide.placeholders[1]
        sub.text = subtitle
        for para in sub.text_frame.paragraphs:
            for run in para.runs:
                run.font.size = Pt(18)
                run.font.color.rgb = RGBColor(*BRAND["dark"])
    return slide


def _add_bullets_slide(prs, doc_title, bullets):
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = doc_title
    title.text_frame.paragraphs[0].runs[0].font.bold = True
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    title.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(*BRAND["title_color"])
    body = slide.placeholders[1].text_frame
    body.clear()
    first = True
    for item in bullets:
        if first:
            para = body.paragraphs[0]
            first = False
        else:
            para = body.add_paragraph()
        para.text = item
        para.level = 0
        para.font.size = Pt(18)
        for run in para.runs:
            run.font.color.rgb = RGBColor(*BRAND["dark"])
    return slide


def _slides_content(kind: str) -> list[tuple[str, list[str]]]:
    if kind == "worker":
        return [
            ("Welcome to Broom Service!", [
                "Thank you for joining our cleaning & co-hosting crew",
                "Today: your app, your job, your pay, and how we act",
                "Short quiz at the end — easy if you pay attention",
            ]),
            ("About Our Company", [
                "We clean & co-host short-term rentals (Airbnb, Vrbo, direct)",
                "Guests fund the cleaning fee — hosts pay nothing out of pocket",
                "We win by doing 5-star turnovers, every time",
            ]),
            ("Your Phone App — Download & Login", [
                "Open https://bizstackperks.com on your phone",
                "Save it to your home screen so it opens like an app",
                "Log in with your email and the PIN the owner set for you",
                "That's it — no store download needed",
            ]),
            ("Your Phone App — Jobs & Map", [
                "Your assigned jobs show with the date, time and address",
                "Tap 'Directions' / the map to navigate to the property",
                "Jobs stay on your screen so you always know your day",
            ]),
            ("Clock In / Clock Out", [
                "When you arrive, the app checks you're at the right property",
                "Clock in when you're inside the work area and start",
                "Clock out when the job is finished (before you leave!)",
                "Clock in AND out on every job — that's what pays you",
            ]),
            ("Paychecks & Pay", [
                "You get paid per job at your rate on your paycheck",
                "Find pay stubs under 'Pay' in the app or on the web portal",
                "Questions about pay? Ask the office—never negotiate in the host's home",
            ]),
            ("Using the Website Portal", [
                "The website (https://bizstackperks.com) shows the same schedule",
                "Log in at /worker-login or through /login with your email + PIN",
                "See your jobs, map, and paychecks in one place",
            ]),
            ("House Rules in Hosts' Homes", [
                "You are a guest in their property — act like it",
                "No smoking, no eating host food, no using beds or bathrooms for personal use",
                "Leave every place exactly as you found it, plus cleaner",
                "Treat the host's stuff the way you'd want yours treated",
            ]),
            ("The 5-Star Turnover Checklist", [
                "Trash out, then all linen exchanged",
                "Bathrooms + kitchen sanitized top to bottom",
                "All floors vacuumed and mopped",
                "Dust everything, restock supplies, stage the space",
                "Photo-verify each room, then mark the job done",
            ]),
            ("Staying Safe at Work", [
                "Use supplies safely; report anything unsafe",
                "Let the office know any medical issue you have before jobs",
                "If you feel unsafe at a property, leave and call the office",
            ]),
            ("Workplace Ethics & Integrity", [
                "Be honest about time, vetting, and what you complete",
                "Don't cut corners to rush — 5-star is the brand",
                "Never share host or guest information with anyone",
            ]),
            ("Sexual Harassment Policy", [
                "Zero tolerance — always, toward anyone, at work",
                "Unwelcome touching, comments, jokes, or advances are prohibited",
                "Reporting is safe and confidential — no retaliation, ever",
            ]),
            ("Reporting Problems & Concerns", [
                "Damage, maintenance, host issues → report right away with a photo",
                "Harassment or anything wrong → tell the owner or use the report channel",
                "Reporting is never punished. Ever.",
            ]),
        ]
    if kind in ("construction", "construction-osha", "osha"):
        return [
            ("Welcome to Buildstack Construction!", [
                "Thank you for joining our licensed general contracting crew",
                "Today: your pay app, job site safety (OSHA-10 baseline), and how we act",
                "Short knowledge test at the end — pass to move to jobs",
            ]),
            ("About Our Company", [
                "Licensed, insured general contractor — resi + commercial",
                "Whole-home renovations, kitchens, baths, drywall, roofing, decks, all trades",
                "We build trust with 5-star work, safety first, every day",
            ]),
            ("Your Pay App — Login & Direct Deposit", [
                "Crew portal at your site login; use your phone + PIN",
                "Punch in/out on every job — that's what pays you",
                "Set up direct deposit (Stripe) so checks land automatically",
                "See hours, timesheets, and pay history under 'Pay'",
            ]),
            ("Clock In / Clock Out", [
                "Clock in when you're on site and ready to work",
                "Clock out when the job is done (before you leave)",
                "Clock in AND out on every job — that's how payroll happens",
                "Report all hours honestly; never punch for someone else",
            ]),
            ("OSHA-10: The Fatal Four", [
                "Falls — the #1 killer on job sites",
                "Struck-by — stay out of equipment swing/reach paths",
                "Caught-in/between — never enter trenches or between moving parts",
                "Electrical — LOTO, test before touching, use GFCI",
            ]),
            ("Fall Protection", [
                "Fall protection required at 6 feet and above",
                "Guardrails, covers, or a harness + anchor — no exceptions",
                "Never work on a roof without tie-off or other approved fall protection",
                "Keep ladders on firm ground: 3 points of contact, don't stand above the top rung",
            ]),
            ("PPE — Every Job, No Exceptions", [
                "Hard hat + eye protection on active job sites",
                "Gloves, steel-toe footwear, hearing protection, hi-vis as the task requires",
                "Check your PPE daily; report damaged gear to the owner",
                "No PPE, no work — period",
            ]),
            ("Hazard Communication & Chemicals", [
                "Read SDS before using any chemical product",
                "Eye/face wash and ventilation when mixing or spraying",
                "Never mix bleach with ammonia or other cleaners (deadly gas)",
                "Use the right chemical for the right task; store safely",
            ]),
            ("Lockout / Tagout (LOTO)", [
                "LOTO prevents accidental energy release — electrical, plumbing, equipment",
                "Only the person who applied the lock removes it",
                "Verify with a tester that circuits are dead before working",
                "Never wedge a switch or bypass a safety guard",
            ]),
            ("Trench & Excavation Safety", [
                "Never enter an unprotected trench — that's caught-in/between risk",
                "Trenches 5 ft+ need shoring, sloping, or a protective system",
                "Keep spoil piles and equipment clear of the trench edge",
                "When in doubt, ask the owner before going in",
            ]),
            ("Silica, Dust & Asbestos Awareness", [
                "Wet cut concrete/masonry + use dust collection",
                "Wear a respirator where silica or dust is heavy",
                "Older builds (pre-1970s) may have asbestos or lead paint — don't disturb it",
                "Report any suspected asbestos/lead to the owner immediately",
            ]),
            ("Workplace Ethics & HR", [
                "Be honest about time, materials, and what you complete",
                "Sexual harassment is zero tolerance — unwanted comments/touching, any time",
                "Report problems confidentially — no retaliation, ever",
                "Never share client or business information with anyone",
            ]),
            ("Your Knowledge Test", [
                "Short test: tape-measure reading, simple math, OSHA basics",
                "Covers framing, drywall, roofing, tile, plumbing basics",
                "Pass it and you're clear to move to jobs",
                "Certificate is recorded in your crew record",
            ]),
        ]
    return [
        ("Welcome to Broom Service", [
            "Short-term rental cleaning + co-hosting, powered by automation",
            "You're the kind of host we love — a property with potential",
            "This deck shows services, pricing, your portal, and growth options",
        ]),
        ("What We Do", [
            "Turnover cleaning, deep cleaning, linen restock, inspections",
            "Digital co-hosting: 24/7 AI assistant, dynamic pricing, review help",
            "Full-service management: cleaning, vendors, multi-channel bookings",
            "All backed by photo-verified quality you can see",
        ]),
        ("Service Pricing", [
            "Turnover Cleaning $120  |  Deep Cleaning $200",
            "Linen Restock $50  |  Inspection $75",
            "Cleaning is GUEST-FUNDED — you pay $0 out of pocket",
            "Co-hosting: 10-15% (digital) or 20-30% (full-service) of gross bookings",
        ]),
        ("How Booking Works", [
            "Guests contact the 24/7 assistant by text or call",
            "Assistant books the slot and sends a secure Stripe payment link",
            "Guest pays at booking — your calendar and ledger update automatically",
        ]),
        ("Your Free Rental Analysis", [
            "Get a real, data-backed earnings report for your property",
            "Home value, market rent, nightly STR estimate, self-managed vs co-hosted",
            "Free via the site — just submit your address",
        ]),
        ("Your Host Portal", [
            "Log in at /host-login to see YOUR properties, bookings and payments",
            "A live map of your portfolio",
            "Only your data — ever",
        ]),
        ("Calendar & Payments", [
            "Bookings sync to your calendar",
            "Paid bookings appear as revenue automatically",
            "Unpaid guests can be re-sent a fresh payment link any time",
        ]),
        ("Growing With Financing", [
            "Need capital to furnish, stage, or convert your property?",
            "Broom Service connects eligible leads with bank/financing partners",
            "Funding is handled directly with the partner bank — we just connect you",
        ]),
        ("What Makes a 5-Star Host", [
            "Fast, friendly guest communication",
            "Immaculate turnover between guests (our job!)",
            "Clear house rules and a clean, stocked property",
        ]),
        ("Support & Next Steps", [
            "Questions? Text/call +1 (757) 846-9275 or email hello@bizstackperks.com",
            "Submit your property for a free analysis",
            "Let's get your listing earning",
        ]),
    ]


def deck_slides(kind: str) -> list[dict]:
    """Deck slides as {title, bullets, voice} — voice is the bot's narration text."""
    parsed = "construction" if kind in ("construction", "construction-osha", "osha") else ("host" if kind != "worker" else "worker")
    phrases = {
        "Welcome to Broom Service!": "Welcome to Broom Service.",
        "Welcome to Broom Service": "Welcome to Broom Service.",
        "About Our Company": "About our company.",
        "Welcome to Buildstack Construction!": "Welcome to Buildstack Construction.",
    }
    out = []
    for title, bullets in _slides_content(parsed):
        voice = phrases.get(title, f"Next up — {title}. ") + " "
        voice += " ".join(b for b in bullets if b)
        out.append({"title": title, "bullets": bullets, "voice": voice})
    return out


def build_deck(kind: str) -> bytes:
    """Build a powerpoint deck. kind: 'worker', 'host', or 'construction'. Returns .pptx bytes."""
    from pptx import Presentation

    is_construction = kind in ("construction", "construction-osha", "osha")
    kind = "construction" if is_construction else ("host" if kind != "worker" else "worker")
    prs = Presentation()
    prs.slide_width = 12192000
    prs.slide_height = 6858000

    if kind == "worker":
        doc_title = "New Worker Orientation"
        subtitle = "Your app, your job, your pay, and how we do it — plus your quick quiz."
    elif kind == "construction":
        doc_title = "Crew Orientation + OSHA-10 Baseline"
        subtitle = "Your pay app, job-site safety (the fatal four), and how we build — plus your knowledge test."
    else:
        doc_title = "Host & Lead Onboarding"
        subtitle = "Welcome to Broom Service — how we help your property earn."

    _add_title_slide(prs, doc_title, subtitle)
    for title, bullets in _slides_content(kind):
        _add_bullets_slide(prs, title, bullets)

    from io import BytesIO
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def grade_quiz(answers: list[int]) -> dict:
    """Grade quiz answers (list of 10 option indexes) -> score/percent + right/wrong."""
    if len(answers) != len(QUIZ):
        raise ValueError("expected 10 answers")
    correct = [i for i, a in enumerate(answers) if a == QUIZ[i]["answer"]]
    pct = round(100 * len(correct) / len(QUIZ))
    passed = pct >= 70
    return {
        "total": len(QUIZ),
        "correct": len(correct),
        "percent": pct,
        "passed": passed,
    }


def grade_osha_quiz(answers) -> dict:
    """Grade the OSHA-10 / trades orientation quiz.

    Accepts a list of {question_id, answer} dicts (answer = option index) or a flat
    list of option indexes. Returns pass/fail, percent, and missed topics.
    """
    if answers and isinstance(answers[0], dict):
        by_id = {}
        for item in answers:
            try:
                qid = int(item.get("question_id") or item.get("qid") or "0")
            except (TypeError, ValueError):
                qid = 0
            try:
                by_id[qid] = int(item.get("answer"))
            except (TypeError, ValueError):
                continue
        flat = [by_id.get(i) for i in range(len(OSHA_QUIZ))]
        answers = flat
    answers = list(answers or [])
    if len(answers) != len(OSHA_QUIZ):
        raise ValueError(f"expected {len(OSHA_QUIZ)} answers")
    correct = [i for i, a in enumerate(answers) if a == OSHA_QUIZ[i]["answer"]]
    missed = sorted({OSHA_QUIZ[i]["topic"] for i in range(len(OSHA_QUIZ)) if i not in correct})
    pct = round(100 * len(correct) / len(OSHA_QUIZ))
    passed = pct >= 70
    return {
        "total": len(OSHA_QUIZ),
        "correct": len(correct),
        "percent": pct,
        "passed": passed,
        "missed_topics": missed,
    }